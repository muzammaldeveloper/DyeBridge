from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Any

from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook


@dataclass
class Recipe:
    serial_code: str
    shade: str
    chemicals: dict[str, float]


def _as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_recipes_from_excel(path: str, start_row: int = 18, end_row: int = 44) -> list[Recipe]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active

    header_row = start_row - 1
    headers: dict[int, str] = {}
    for col in range(3, sheet.max_column + 1):
        header = sheet.cell(row=header_row, column=col).value
        if header:
            headers[col] = str(header).strip()

    recipes: list[Recipe] = []
    for row in range(start_row, end_row + 1):
        serial = sheet.cell(row=row, column=1).value
        shade = sheet.cell(row=row, column=2).value
        if not serial or not shade:
            continue

        chemicals: dict[str, float] = {}
        for col in range(3, sheet.max_column + 1):
            percent = _as_float(sheet.cell(row=row, column=col).value)
            if percent is None:
                continue

            name = headers.get(col)
            if name:
                chemicals[name] = percent

        if not chemicals:
            chem_name = sheet.cell(row=row, column=3).value
            chem_pct = _as_float(sheet.cell(row=row, column=4).value)
            if chem_name and chem_pct is not None:
                chemicals[str(chem_name).strip()] = chem_pct

        recipes.append(
            Recipe(
                serial_code=str(serial).strip(),
                shade=str(shade).strip(),
                chemicals=chemicals,
            )
        )

    return recipes


def compute_quantities(chemicals: dict[str, float], greigh_weight_kg: float) -> list[dict[str, float | str]]:
    result: list[dict[str, float | str]] = []
    for name, pct in chemicals.items():
        qty_kg = (greigh_weight_kg * pct) / 100
        qty_g = qty_kg * 1000
        qty_mg = qty_g * 1000
        result.append(
            {
                "chemical": name,
                "percentage": round(pct, 4),
                "kg": round(qty_kg, 6),
                "g": round(qty_g, 3),
                "mg": round(qty_mg, 2),
            }
        )
    return result


def create_app() -> Flask:
    app = Flask(__name__)
    app.config["RECIPES_EXCEL_PATH"] = os.environ.get(
        "RECIPES_EXCEL_PATH",
        os.path.join(app.root_path, "data", "shade_recipes.xlsx"),
    )

    @app.get("/")
    def index() -> str:
        return render_template("index.html")

    @app.get("/api/recipes")
    def list_recipes():
        excel_path = request.args.get("excel_path") or app.config["RECIPES_EXCEL_PATH"]
        if not os.path.exists(excel_path):
            return jsonify({"error": f"Excel file not found: {excel_path}"}), 404

        recipes = load_recipes_from_excel(excel_path)
        return jsonify(
            [
                {
                    "serial_code": recipe.serial_code,
                    "shade": recipe.shade,
                    "chemicals": recipe.chemicals,
                }
                for recipe in recipes
            ]
        )

    @app.post("/api/calculate")
    def calculate_recipe():
        payload = request.get_json(silent=True) or {}
        chemicals = payload.get("chemicals") or {}
        greigh_weight = _as_float(payload.get("greigh_weight_kg"))
        if greigh_weight is None or greigh_weight <= 0:
            return jsonify({"error": "greigh_weight_kg must be a positive number"}), 400

        return jsonify(
            {
                "machine_no": payload.get("machine_no", ""),
                "cones": payload.get("cones", ""),
                "lot_no": payload.get("lot_no", ""),
                "po_no": payload.get("po_no", ""),
                "mpo_no": payload.get("mpo_no", ""),
                "greigh_weight_kg": greigh_weight,
                "items": compute_quantities(chemicals, greigh_weight),
            }
        )

    return app


app = create_app()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
