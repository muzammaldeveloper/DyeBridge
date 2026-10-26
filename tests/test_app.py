from pathlib import Path
import sys

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app, load_recipes_from_excel


def _write_sample_excel(path: str) -> None:
    wb = Workbook()
    ws = wb.active

    ws.cell(row=17, column=1, value="Recipe")
    ws.cell(row=17, column=2, value="Shade")
    ws.cell(row=17, column=3, value="Blue Dye")
    ws.cell(row=17, column=4, value="Fixing Chemical")

    ws.cell(row=18, column=1, value="RCP-001")
    ws.cell(row=18, column=2, value="Navy")
    ws.cell(row=18, column=3, value=1.5)
    ws.cell(row=18, column=4, value=0.25)

    wb.save(path)


def test_excel_parser_reads_rows_18_to_44(tmp_path):
    excel_file = tmp_path / "recipes.xlsx"
    _write_sample_excel(str(excel_file))

    recipes = load_recipes_from_excel(str(excel_file))

    assert len(recipes) == 1
    assert recipes[0].serial_code == "RCP-001"
    assert recipes[0].shade == "Navy"
    assert recipes[0].chemicals == {"Blue Dye": 1.5, "Fixing Chemical": 0.25}


def test_calculate_api_returns_kg_g_mg():
    app = create_app()
    client = app.test_client()

    response = client.post(
        "/api/calculate",
        json={
            "greigh_weight_kg": 100,
            "chemicals": {"Blue Dye": 1.5},
        },
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["items"][0]["kg"] == 1.5
    assert payload["items"][0]["g"] == 1500.0
    assert payload["items"][0]["mg"] == 1500000.0
